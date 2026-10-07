# -*- coding: utf-8 -*-
"""Sinh sáu DOCX từ một nguồn nội dung, có bìa, mở đầu, cảm ơn, cam đoan, mục lục."""
from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import trich_xuat

WIDTHS={
    "so_do_khoi_1_truc.png":15.5,"so_do_khoi_2_truc.png":15.5,
    "so_do_ket_noi_1_truc.png":16.0,"so_do_ket_noi_2_truc.png":16.0,
    "bo_tri_4_ldr.png":15.8,"so_do_3_phuong_phap.png":15.5,
    "duong_di_mat_troi_1_truc.png":15.0,"hoat_dong_hybrid_1_truc.png":15.5,
    "hoat_dong_hybrid_2_truc.png":15.5,"so_sanh_nang_luong_1_truc.png":15.5,
    "so_sanh_nang_luong_2_truc.png":15.5,
    "luu_do_tong_quat_1_truc.png":11.5,"luu_do_tong_quat_2_truc.png":11.5,
    "luu_do_thien_van.png":11.3,"luu_do_doc_adc.png":11.3,
    "luu_do_dieu_khien_motor.png":11.3,"luu_do_lcd.png":11.3,
}


def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement("w:shd");shd.set(qn("w:fill"),fill);tcPr.append(shd)


def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr();el=OxmlElement("w:tblHeader");el.set(qn("w:val"),"true");trPr.append(el)


def set_field(run,instruction):
    begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText");instr.set(qn("xml:space"),"preserve");instr.text=instruction
    separate=OxmlElement("w:fldChar");separate.set(qn("w:fldCharType"),"separate")
    text=OxmlElement("w:t");text.text="Cập nhật trường trong Word để hiện số trang/mục lục."
    end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
    run._r.extend([begin,instr,separate,text,end])


def setup_doc(doc,item):
    section=doc.sections[0]
    section.top_margin=Cm(2.2);section.bottom_margin=Cm(2.0)
    section.left_margin=Cm(2.6);section.right_margin=Cm(2.2)
    section.header_distance=Cm(1.1);section.footer_distance=Cm(1.0)
    styles=doc.styles
    normal=styles["Normal"]
    normal.font.name="Times New Roman";normal.font.size=Pt(12);normal.font.color.rgb=RGBColor(20,20,20)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
    normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.first_line_indent=Cm(0.95)
    for name,size,color in (("Heading 1",15,"1F3864"),("Heading 2",13,"1F3864"),("Heading 3",12,"1F3864")):
        s=styles[name];s.font.name="Times New Roman";s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor.from_string(color)
        s._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
        s.paragraph_format.keep_with_next=True;s.paragraph_format.first_line_indent=Cm(0)
        s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(5)
    # Header and page number.
    hp=section.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.first_line_indent=Cm(0)
    r=hp.add_run(f"{item['quyen']}  |  {item['subtitle']}");r.font.name="Times New Roman";r.font.size=Pt(9);r.font.color.rgb=RGBColor(90,90,90)
    fp=section.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=Cm(0)
    rr=fp.add_run("Trang ");rr.font.name="Times New Roman";rr.font.size=Pt(9)
    set_field(fp.add_run(),"PAGE")
    # Ask Word to refresh generated TOC/page fields on opening.
    settings=doc.settings.element
    update=OxmlElement("w:updateFields");update.set(qn("w:val"),"true");settings.append(update)
    doc.core_properties.title=f"{item['quyen']} – {item['subtitle']}"
    doc.core_properties.subject=trich_xuat.DE_TAI
    doc.core_properties.author="Nhóm thực hiện đồ án"


def centered(doc,text,size,bold=False,color="1F3864",after=8):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(after)
    r=p.add_run(text);r.bold=bold;r.font.name="Times New Roman";r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color)
    return p


def add_frontmatter(doc,item):
    for _ in range(3):doc.add_paragraph()
    centered(doc,"BÁO CÁO ĐỒ ÁN",17,True,after=8)
    centered(doc,trich_xuat.DE_TAI,19,True,after=18)
    centered(doc,item["quyen"],15,True,after=10)
    centered(doc,item["volume_label"].upper()+" – CẤU HÌNH "+item["axis_label"].upper(),13,True,after=18)
    centered(doc,"Phương pháp hybrid: thiên văn định vị thô, LDR hiệu chỉnh khi đủ sáng; ban đêm/thiếu sáng giữ vị trí.",11,False,"333333",after=28)
    centered(doc,"Ngày biên soạn: 07/10/2026",11,False,"333333",after=4)
    centered(doc,"",11)
    doc.add_page_break()

    doc.add_heading("LỜI MỞ ĐẦU",level=1)
    doc.add_paragraph("Báo cáo trình bày phần việc được phân công trong đồ án hệ thống bám nắng mặt trời. Sáu quyển được chia theo cấu hình một trục/hai trục và theo nhóm Nghiên cứu, Chế tạo, Lập trình để tách rõ cơ sở tính toán, giao diện phần cứng và mã điều khiển.")
    doc.add_paragraph("Các thông số thống nhất được quản lý trong Bảng thông số chung đặt trước bộ báo cáo. Số liệu trong đồ án được ghi đúng bản chất: kết quả tính toán hoặc mô phỏng không được trình bày như phép đo; các công việc chưa thực hiện được nêu kèm kế hoạch dự kiến.")
    doc.add_page_break()

    doc.add_heading("LỜI CẢM ƠN",level=1)
    doc.add_paragraph("Nhóm thực hiện trân trọng cảm ơn giảng viên hướng dẫn, thầy cô và các bạn đã hỗ trợ định hướng, góp ý và tạo điều kiện trong quá trình xây dựng đề tài. Những nhận xét chuyên môn giúp nhóm thống nhất lại quy ước góc, giới hạn hành trình, phân công nội dung và cách trình bày bằng chứng.")
    doc.add_paragraph("Nhóm xin tiếp thu các góp ý còn lại để hoàn thiện bản vẽ, kiểm thử phần cứng và hiệu chuẩn trước khi nghiệm thu.")
    doc.add_page_break()

    doc.add_heading("LỜI CAM ĐOAN",level=1)
    doc.add_paragraph("Nhóm cam đoan nội dung được trình bày theo đúng phạm vi quyển báo cáo và nguồn tham khảo đã dẫn. Kết quả tính toán/mô phỏng được phân biệt với kết quả đo. Báo cáo không khẳng định đã lắp ráp, hiệu chuẩn, đo dòng/mô-men hoặc thử nghiệm phần cứng khi chưa có biên bản xác nhận; các việc đó được ghi là chưa thực hiện và có thời hạn dự kiến.")
    doc.add_paragraph("Các thông số cơ khí, tải gió, công suất nhãn motor và giới hạn góc chưa được xác nhận được đánh dấu là giả thiết thiết kế, không phải kết quả thực nghiệm.")
    doc.add_page_break()

    doc.add_heading("MỤC LỤC",level=1)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
    set_field(p.add_run(),'TOC \\o "1-2" \\h \\z \\u')
    p=doc.add_paragraph("Nếu mục lục chưa cập nhật, chọn Update Field / Update entire table trong Word.")
    p.paragraph_format.first_line_indent=Cm(0)
    p.runs[0].italic=True;p.runs[0].font.size=Pt(9)


def add_block(doc,block):
    kind=block[0]
    if kind in ("h1","h2","h3"):
        p=doc.add_heading(block[1],level={"h1":1,"h2":2,"h3":3}[kind])
        p.paragraph_format.first_line_indent=Cm(0)
        if kind=="h1" and "TÀI LIỆU THAM KHẢO" not in block[1]:p.paragraph_format.page_break_before=True
    elif kind=="p":
        p=doc.add_paragraph(block[1]);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    elif kind=="b":
        p=doc.add_paragraph(block[1],style="List Bullet");p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent=Cm(0)
    elif kind=="eq":
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        r=p.add_run(block[1]);r.italic=True;r.font.size=Pt(11)
    elif kind=="code":
        for i,line in enumerate(block[1].splitlines()):
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(.4)
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_together=True
            r=p.add_run(line if line else " ");r.font.name="Consolas";r.font.size=Pt(8);r.font.color.rgb=RGBColor(31,51,80)
    elif kind=="img":
        full=os.path.join(ROOT,block[1])
        if not os.path.exists(full):raise FileNotFoundError(full)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        p.add_run().add_picture(full,width=Cm(WIDTHS.get(os.path.basename(full),15.0)))
        if block[2]:
            c=doc.add_paragraph(block[2],style="Caption");c.alignment=WD_ALIGN_PARAGRAPH.CENTER;c.paragraph_format.first_line_indent=Cm(0)
    elif kind=="tbl":
        rows,caption=block[1],block[2]
        if caption:
            p=doc.add_paragraph(caption,style="Caption");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        table=doc.add_table(rows=len(rows),cols=len(rows[0]));table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=True
        for i,row in enumerate(rows):
            if i==0:set_repeat_table_header(table.rows[i])
            for j,value in enumerate(row):
                cell=table.cell(i,j);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                if i==0:set_cell_shading(cell,"D9E5F3")
                p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
                r=p.add_run(str(value));r.font.name="Times New Roman";r.font.size=Pt(8.5);r.bold=(i==0)
        spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.first_line_indent=Cm(0)


def build_one(item):
    doc=Document();setup_doc(doc,item);add_frontmatter(doc,item)
    for block in item["blocks"]:add_block(doc,block)
    doc.save(item["docx"])
    print("đã tạo",os.path.basename(item["docx"]),"|",len(item["blocks"]),"khối")


def main():
    for item in trich_xuat.danh_sach_bao_cao():build_one(item)


if __name__=="__main__":main()
