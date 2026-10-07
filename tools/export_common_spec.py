# -*- coding: utf-8 -*-
"""Xuất bảng thông số gốc dùng chung ra Markdown, DOCX và một trang PDF."""
from __future__ import annotations

import os
import sys
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4,landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle

ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import thong_so_chung as spec


def shade(cell,color):
    pr=cell._tc.get_or_add_tcPr();el=OxmlElement("w:shd");el.set(qn("w:fill"),color);pr.append(el)


def make_docx(rows,path):
    doc=Document();s=doc.sections[0];s.orientation=WD_ORIENT.LANDSCAPE
    s.page_width=Cm(29.7);s.page_height=Cm(21.0)
    s.top_margin=Cm(1.0);s.bottom_margin=Cm(1.0);s.left_margin=Cm(1.1);s.right_margin=Cm(1.1)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
    r=p.add_run("BẢNG THÔNG SỐ CHUNG – HỆ THỐNG BÁM NẮNG");r.bold=True;r.font.size=Pt(14);r.font.color.rgb=RGBColor(31,56,100)
    p=doc.add_paragraph(f"Phiên bản {spec.SPEC_VERSION} · {spec.SPEC_DATE} · nguồn chuẩn thống nhất cho sáu quyển");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(5);p.runs[0].font.size=Pt(8.5)
    t=doc.add_table(rows=1,cols=2);t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    t.columns[0].width=Cm(4.5);t.columns[1].width=Cm(22.8)
    for cell,text in zip(t.rows[0].cells,("Hạng mục","Quy ước / giá trị thống nhất")):
        cell.text=text;shade(cell,"D9E5F3")
        for r in cell.paragraphs[0].runs:r.bold=True;r.font.size=Pt(8)
    for key,value in rows:
        cells=t.add_row().cells;cells[0].width=Cm(4.5);cells[1].width=Cm(22.8)
        for cell,text in zip(cells,(key,value)):
            cell.text=text;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.size=Pt(7.6);r.font.name="Times New Roman"
    p=doc.add_paragraph("Tải gió sơ bộ: q=0,5ρV²; F=q·Cd·A; M_az=1,5·F·0,30; M_tilt=1,5·(F·0,20+3). Mô-men 20 N·m là giả thiết, chưa xác minh; không phải mô-men liên tục.")
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(0);p.runs[0].font.size=Pt(7.2)
    p=doc.add_paragraph("Chưa xác nhận lắp ráp/hiệu chuẩn/đo dòng-mô-men/thử ngoài trời/phần cứng ESP32. Dự kiến hoàn tất trước 10/10/2026.")
    p.paragraph_format.space_after=Pt(0);p.runs[0].font.size=Pt(7.2);p.runs[0].bold=True
    doc.core_properties.title="Bảng thông số chung hệ thống bám nắng"
    doc.core_properties.author="Nhóm thực hiện đồ án"
    doc.save(path)


def make_pdf(rows,path):
    fontdir=os.path.join(__import__("matplotlib").get_data_path(),"fonts","ttf")
    pdfmetrics.registerFont(TTFont("SpecSans",os.path.join(fontdir,"DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("SpecSansB",os.path.join(fontdir,"DejaVuSans-Bold.ttf")))
    styles={
        "cell":ParagraphStyle("cell",fontName="SpecSans",fontSize=6.7,leading=7.7,textColor=colors.HexColor("#17202a")),
        "head":ParagraphStyle("head",fontName="SpecSansB",fontSize=7,leading=8,textColor=colors.HexColor("#17202a")),
        "title":ParagraphStyle("title",fontName="SpecSansB",fontSize=13,leading=15,alignment=1,textColor=colors.HexColor("#1f3864")),
        "sub":ParagraphStyle("sub",fontName="SpecSans",fontSize=7.2,leading=9,alignment=1),
        "foot":ParagraphStyle("foot",fontName="SpecSans",fontSize=6.6,leading=7.8),
    }
    data=[[Paragraph("Hạng mục",styles["head"]),Paragraph("Quy ước / giá trị thống nhất",styles["head"])]]
    for key,value in rows:
        data.append([Paragraph(escape(key),styles["head"]),Paragraph(escape(value),styles["cell"])])
    tbl=Table(data,colWidths=[48*mm,222*mm],repeatRows=1,hAlign="CENTER")
    tbl.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#667085")),
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#d9e5f3")),
        ("BACKGROUND",(0,1),(0,-1),colors.HexColor("#f0f3f7")),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
        ("TOPPADDING",(0,0),(-1,-1),2.2),("BOTTOMPADDING",(0,0),(-1,-1),2.2),
    ]))
    story=[Paragraph("BẢNG THÔNG SỐ CHUNG – HỆ THỐNG BÁM NẮNG",styles["title"]),
           Paragraph(f"Phiên bản {spec.SPEC_VERSION} · nguồn chuẩn cho sáu quyển · {spec.SPEC_DATE}",styles["sub"]),
           Spacer(1,3*mm),tbl,Spacer(1,2*mm),
           Paragraph("Tải gió sơ bộ: q=0,5ρV²; F=q·Cd·A; M_az=1,5·F·0,30; M_tilt=1,5·(F·0,20+3). 20 N·m là giả thiết chưa xác minh, không phải mô-men liên tục.",styles["foot"]),
           Paragraph("Trạng thái: chưa xác nhận lắp ráp, hiệu chuẩn, đo dòng/mô-men, thử ngoài trời hoặc chạy trên phần cứng. Dự kiến hoàn tất trước 10/10/2026.",styles["foot"])]
    doc=SimpleDocTemplate(path,pagesize=landscape(A4),leftMargin=12*mm,rightMargin=12*mm,topMargin=8*mm,bottomMargin=8*mm,title="Bảng thông số chung")
    doc.build(story)


def main():
    rows=spec.common_spec_rows()
    md=spec.shared_spec_markdown()
    for filename in ("Bang_thong_so_chung.md","Bang_thong_so_chung.docx","Bang_thong_so_chung.pdf"):
        path=os.path.join(ROOT,filename)
        if filename.endswith(".md"):
            with open(path,"w",encoding="utf-8") as f:f.write(md)
        elif filename.endswith(".docx"):make_docx(rows,path)
        else:make_pdf(rows,path)
        print("đã xuất",filename)


if __name__=="__main__":main()
